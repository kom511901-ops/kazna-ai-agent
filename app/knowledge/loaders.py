"""Асинхронная загрузка PDF, DOCX, HTML и текстовых файлов."""

import asyncio
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as WordDocument
from pypdf import PdfReader

SUPPORTED_SUFFIXES = frozenset({".pdf", ".docx", ".html", ".htm", ".txt"})


@dataclass(frozen=True, slots=True)
class LoadedDocument:
    """Извлечённый чистый текст и путь исходного файла."""

    source_path: Path
    text: str


async def load_document(path: Path) -> LoadedDocument:
    """Загружает поддерживаемый файл и возвращает очищенный текст."""
    normalized_path = path.expanduser().resolve()
    suffix = normalized_path.suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Неподдерживаемый формат документа: {suffix or 'без расширения'}")

    content = await asyncio.to_thread(normalized_path.read_bytes)
    extracted = await asyncio.to_thread(_extract_text, content, suffix)
    cleaned = clean_text(extracted)
    if not cleaned:
        raise ValueError(f"В документе не найден текст: {normalized_path}")
    return LoadedDocument(source_path=normalized_path, text=cleaned)


async def discover_documents(source: Path) -> list[Path]:
    """Возвращает файлы поддерживаемых форматов из каталога или путь к одному файлу."""
    normalized_path = source.expanduser().resolve()

    def find_files() -> list[Path]:
        if normalized_path.is_file():
            if normalized_path.suffix.casefold() not in SUPPORTED_SUFFIXES:
                raise ValueError(f"Неподдерживаемый формат документа: {normalized_path.suffix}")
            return [normalized_path]
        if not normalized_path.is_dir():
            raise FileNotFoundError(f"Путь к источникам не найден: {normalized_path}")
        return sorted(
            path
            for path in normalized_path.rglob("*")
            if path.is_file() and path.suffix.casefold() in SUPPORTED_SUFFIXES
        )

    return await asyncio.to_thread(find_files)


def clean_text(value: str) -> str:
    """Нормализует пробелы и строки, сохраняя границы абзацев."""
    normalized = value.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"[\t\f\v ]+", " ", normalized)
    normalized = re.sub(r" *\n *", "\n", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _extract_text(content: bytes, suffix: str) -> str:
    """Выбирает парсер по расширению; выполняется во вспомогательном потоке."""
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(content))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        document = WordDocument(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_rows = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join([*paragraphs, *table_rows])
    if suffix in {".html", ".htm"}:
        soup = BeautifulSoup(content, "html.parser")
        for element in soup(["script", "style", "noscript"]):
            element.decompose()
        return soup.get_text("\n")
    if suffix == ".txt":
        try:
            return content.decode("utf-8-sig")
        except UnicodeDecodeError:
            return content.decode("cp1251")
    raise ValueError(f"Неподдерживаемый формат документа: {suffix}")
